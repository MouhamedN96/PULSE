# PULSE: Voice-Curated Local Social Discovery
# Final Project Paper — Rewrite 
# Mouhamed Ndiaye & Itayi Penda

"""
This script generates a rewritten version of the PULSE paper
with natural, human-sounding prose while preserving all technical content.
"""

import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# ── Style setup ──
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
    return h

def para(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    return p

# ── Title Block ──
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('PULSE: Voice-Curated Social Activity Discovery\nUsing Rust, Flutter, and Generative AI')
run.bold = True
run.font.size = Pt(16)
run.font.name = 'Times New Roman'

authors = doc.add_paragraph()
authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = authors.add_run('Mouhamed Ndiaye & Itayi Penda')
run.font.size = Pt(12)
run.font.name = 'Times New Roman'

doc.add_paragraph()

# ── Abstract ──
heading('Abstract', level=1)
para(
    'Finding something to do in a city shouldn\u2019t require scrolling through five apps. '
    'PULSE is a progressive web app that takes a single voice or text query\u2014something like '
    '\u201cfun rooftop bar for a birthday\u201d\u2014and returns a ranked set of venue cards with AI-written '
    'summaries, distance estimates, and real-time data pulled from Google Places, Yelp, and '
    'Eventbrite. The backend is written in Rust (Axum), which keeps memory usage low and '
    'response times fast. A Redis cache layer sits in front of external API calls so that '
    'repeated nearby searches resolve in under 80\u2009ms. On the frontend, Flutter renders the '
    'results as swipeable cards that feel native on both mobile and desktop browsers. '
    'We ran seven experiments covering latency, caching behavior, AI summary quality, voice '
    'intent parsing, distance accuracy, social-graph query speed, and analytics readiness. '
    'The results confirm that the architecture is viable for prototype-scale deployment '
    'and outline a clear path toward production.'
)

# ── Introduction ──
heading('Introduction', level=1)
para(
    'Most local discovery tools fall into one of two camps: keyword-driven search engines '
    'that expect you to already know what you want, or social feeds that bury useful '
    'recommendations under unrelated content. Neither works well when someone just wants '
    'to figure out what to do tonight.'
)
para(
    'PULSE tries to close that gap. A user opens the web app, grants location access, and '
    'speaks or types a natural-language request. The system parses the intent, pulls live '
    'venue data from multiple APIs, computes approximate distances, asks a large language '
    'model to write short summaries, and returns everything as a set of cards the user can '
    'swipe through. The whole round-trip targets under two seconds for a fresh query and '
    'well under one second for a cached repeat.'
)
para(
    'We chose this problem because both of us have lived the frustration of moving to a new '
    'city and not knowing where to start. Yelp gives you a list, but it doesn\u2019t explain why '
    'a place fits your mood. Google Maps gives you directions, but it doesn\u2019t know you\u2019re '
    'looking for somewhere chill to work. We wanted a tool that understands intent, not just '
    'keywords.'
)
para(
    'The project sits at the intersection of recommender systems, location-based services, '
    'and applied NLP. It also touches on data analytics\u2014every request generates logs that '
    'can feed dashboards for cache-hit rates, popular categories, and response-time '
    'distributions\u2014and it includes a social-graph schema in PostgreSQL so that future '
    'versions can weight recommendations by what your friends have saved or visited.'
)
para(
    'Potential applications range from tourism and campus guides to hotel concierge kiosks '
    'and local event platforms. At scale, the same log data could support trend detection '
    'across cities and time-of-day demand forecasting, which is where the project connects '
    'to big-data analytics.'
)

# ── Related Work ──
heading('Related Work', level=1)
para(
    'Recommender systems have been studied for decades, mostly in the context of movies, '
    'music, or e-commerce (Resnick & Varian, 1997; Ricci et al., 2022). The core challenge '
    'is the same\u2014help users pick from too many options\u2014but the venue-discovery problem '
    'adds two twists: the catalog changes daily (a restaurant might close early or a pop-up '
    'might appear), and geography matters a lot. A five-star restaurant across the river is '
    'useless if you\u2019re hungry right now.'
)
para(
    'Point-of-interest (POI) recommendation work (Ye et al., 2011; Liu et al., 2017) '
    'addresses the geography part. Those papers show that distance should be a first-class '
    'ranking signal, not an afterthought. PULSE follows that idea by computing Haversine '
    'distance server-side so that every result comes pre-sorted by proximity.'
)
para(
    'Social recommendation is another thread. Koren et al. (2009) and Rendle et al. (2009) '
    'showed that collaborative signals\u2014what people similar to you have liked\u2014can beat '
    'content-based filtering. We\u2019ve designed PULSE\u2019s PostgreSQL schema to store follow '
    'relationships, saved places, and visit history, but the current prototype doesn\u2019t yet '
    'train a model on those signals. That\u2019s planned for a future iteration.'
)
para(
    'On the NLP side, transformer-based models (Devlin et al., 2019; Reimers & Gurevych, '
    '2019) have made it feasible to parse fuzzy queries like \u201csomething chill for two\u201d into '
    'structured intent. PULSE delegates this to Gemini 2.0 Flash, which handles intent '
    'extraction and summary generation in a single call. This avoids the overhead of '
    'training and hosting a custom model.'
)
para(
    'From a systems perspective, the project draws on the principles in Kleppmann (2017): '
    'cache aggressively, keep the hot path fast, and design for observability from the start.'
)

# ── Dataset ──
heading('Dataset', level=1)
para(
    'PULSE doesn\u2019t rely on a static benchmark dataset. The whole point of the system is '
    'that venue data is live\u2014ratings change, hours get updated, new places open. Each user '
    'request triggers a fresh data-collection mini-pipeline: the backend calls Google Places '
    'Text Search (and optionally Yelp Fusion and Eventbrite), cleans the results, and '
    'enriches them with distance and AI summaries before responding.'
)
para(
    'Google Places provides the backbone: venue name, rating, review count, hours, photos, '
    'address, and coordinates. Yelp adds richer category tags and price-level data, which '
    'helps with queries like \u201ccheap eats.\u201d Eventbrite brings in time-bound activities\u2014'
    'concerts, workshops, festivals\u2014that static venue data can\u2019t cover.'
)
para(
    'Data quality wasn\u2019t always clean. Some records came back without hours or photos. '
    'GPS accuracy dropped indoors, which occasionally distorted distance rankings. We '
    'handled these with safe defaults, validation checks, and fallback text. The caching '
    'layer also smooths things out: once a clean result is in Redis, subsequent users in '
    'the same area get the validated version instantly.'
)
para(
    'To give a sense of what the pipeline handles, here are three representative queries '
    'and what the system returns:'
)
para(
    '\u2022 "Quiet cafe near me for work" \u2192 Cards favoring calm cafes at short distances, '
    'with a summary explaining why each spot is good for getting things done.\n'
    '\u2022 "Fun rooftop bar for a birthday group" \u2192 Bars and lounges sorted by proximity, '
    'with summaries highlighting group-friendly vibes.\n'
    '\u2022 "Something cheap to do tonight" \u2192 Mixed venue types; Gemini interprets the vague '
    'intent and surfaces low-cost social options.'
)

# ── Data Collection ──
heading('Data Collection', level=1)
para(
    'The collection pipeline kicks off in the Flutter frontend. The user opens the app, '
    'grants location access, and enters a query (typed or spoken). The frontend sends '
    'coordinates, radius, and query text to the Rust backend over HTTPS.'
)
para(
    'On the backend, the first thing that happens is cache lookup. We build a cache key '
    'from the rounded latitude/longitude, radius, category, and intent string. If Redis '
    'has a fresh hit, we return it immediately\u2014no external calls needed. If not, the '
    'backend fans out to Google Places (and optionally Yelp and Eventbrite), receives raw '
    'venue JSON, validates and cleans each record, computes Haversine distance from the '
    'user\u2019s coordinates, sends a batch of venue fields to Gemini for summarization, '
    'stitches the summaries back onto the venue cards, writes the full result to Redis, '
    'and returns JSON to the frontend.'
)
para(
    'The main pain points during collection were: incomplete venue records (missing hours, '
    'null photo references), indoor GPS drift affecting distance calculations, API rate '
    'limits making caching essential rather than optional, and occasional Gemini cold-start '
    'delays on the first request of a session.'
)

# ── Approach and Contributions ──
heading('Approach and Original Contributions', level=1)
para(
    'PULSE is organized in four layers, each handling a distinct concern:'
)
para(
    'Layer 1 \u2014 Flutter PWA. The frontend captures location, accepts voice and text input, '
    'renders filter chips, and displays result cards. Flutter was chosen because it delivers '
    'a native-feeling mobile UI through the browser, which lowers the barrier for first-time '
    'users who don\u2019t want to install an app.'
)
para(
    'Layer 2 \u2014 Rust/Axum Backend. This is the orchestration layer: request validation, '
    'cache-key construction, external API calls, distance computation, and response '
    'formatting. We picked Rust for its memory safety and raw performance. Axum gives us '
    'clean, type-safe route definitions without the boilerplate of heavier frameworks.'
)
para(
    'Layer 3 \u2014 Data + AI. Google Places provides live venue data; Gemini 2.0 Flash handles '
    'intent parsing and summary writing; Haversine runs in the backend to rank by proximity; '
    'Redis caches recent results; PostgreSQL stores the social graph and user records. '
    'Together, these form a request-time recommendation pipeline.'
)
para(
    'Layer 4 \u2014 Analytics. Every request generates structured logs: response time, cache '
    'status, intent string, result count, selected card, and failure reason (if any). These '
    'feed into data-analytics dashboards for system tuning. At scale, the same fields '
    'support big-data analytics\u2014trend detection across cities, demand forecasting, and '
    'learned ranking models.'
)
para(
    'In terms of original contributions, we see six:'
)
para(
    '1. Pairing a Flutter PWA with a Rust backend for local discovery\u2014an unusual stack '
    'choice that gives us both a polished mobile feel and low server overhead.\n'
    '2. A server-side enrichment pipeline that cleans raw venue records, attaches distance, '
    'and generates plain-language summaries before the data ever reaches the client.\n'
    '3. Using Redis not just as a performance optimization but as a core UX feature: nearby '
    'users asking similar questions get near-instant results.\n'
    '4. A PostgreSQL social-graph schema that\u2019s ready for friend-based ranking even though '
    'we haven\u2019t trained a model on it yet.\n'
    '5. An analytics-ready logging design that captures the fields needed to measure '
    'what users search for, where the system is slow, and how ranking could improve.\n'
    '6. A big-data analytics roadmap for using large-scale request logs to detect trends, '
    'improve cache placement, and forecast traffic.'
)

# ── Experiments and Results ──
heading('Experiments and Results', level=1)
para(
    'We ran seven experiments. The goal wasn\u2019t to hit production benchmarks\u2014it was to '
    'validate that the architecture makes sense for a real local-discovery flow.'
)

para(
    'Experiment 1: Geo-query latency. We compared fresh round-trips (Google Places + Gemini) '
    'against Redis cache hits. Fresh calls took noticeably longer; cache hits came back in '
    'under 80\u2009ms. The takeaway is straightforward: caching isn\u2019t optional for a mobile-grade '
    'experience.'
)
para(
    'Experiment 2: Cache analytics. We logged cache status, result count, and response time '
    'across repeated nearby searches. The speed gains were clear and consistent. These logs '
    'also turned out to be useful for tuning cache TTLs and estimating API cost savings.'
)
para(
    'Experiment 3: AI summary quality. We showed users raw venue cards alongside Gemini-enriched '
    'cards and asked which set was easier to act on. The enriched cards won\u2014users said the '
    'short summaries made it faster to decide without reading every field.'
)
para(
    'Experiment 4: Voice intent parsing. We fed natural-language queries to Gemini and compared '
    'the parsed intent against hand-labeled ground truth. Common phrases parsed well; unusual '
    'phrasing occasionally missed. The app includes an \u201cedit filters\u201d step as a safety net.'
)
para(
    'Experiment 5: Distance ranking. We compared Haversine ordering against a route-service '
    'baseline. Haversine was accurate enough for \u201cwhich places are closest?\u201d but obviously '
    'can\u2019t account for traffic or one-way streets. Our compromise: use Haversine for initial '
    'ranking, call a route service only when the user taps \u201cget directions.\u201d'
)
para(
    'Experiment 6: Social-graph speed. We measured PostgreSQL join performance for friend-discovery '
    'queries on a small test graph. Queries were fast at that scale, but we\u2019d need proper indexes '
    'and query planning before scaling to thousands of users.'
)
para(
    'Experiment 7: Big-data readiness. We reviewed the log schema to confirm it could grow from '
    'prototype-scale records into a dataset suitable for trend mining, capacity planning, and '
    'learned ranking. The fields are there; the infrastructure to process them at scale is not\u2014'
    'that\u2019s future work.'
)

# Discussion
heading('Discussion', level=2)
para(
    'The latency numbers tell a clear story: caching is the single most impactful thing we built. '
    'A fresh request depends entirely on external services, and there\u2019s only so much we can do '
    'about Google\u2019s response time. But a warm cache hit bypasses all of that and returns in '
    'under 80\u2009ms, which is fast enough that the UI feels instant. Redis also cuts API costs, '
    'which matters more than we expected once we started hitting rate limits during testing.'
)
para(
    'The summary experiment confirmed our intuition that raw data isn\u2019t enough. Users don\u2019t '
    'want to mentally diff five venue cards\u2014they want someone to tell them \u201cthis one\u2019s '
    'closest and has great reviews.\u201d Gemini does that reasonably well, though we learned to '
    'keep the prompt grounded in the actual venue fields. When we let the model speculate, '
    'it occasionally hallucinated details.'
)
para(
    'Voice input worked better than we expected for common queries but worse for anything '
    'unusual. \u201cBest pizza near me\u201d parsed perfectly; \u201csomewhere to kill time before my '
    'flight\u201d was hit-or-miss. The editable-filter fallback is important.'
)
para(
    'Haversine distance is a good-enough heuristic. It\u2019s not road distance and it doesn\u2019t '
    'know about bridges or tunnels, but it runs inside the Rust backend with zero external '
    'calls, which keeps the hot path fast. A route service is better for turn-by-turn '
    'navigation, so we defer that to the point where the user has already picked a venue.'
)
para(
    'The social-graph queries worked at test scale, but we\u2019re under no illusion that a '
    'ten-user graph proves anything about production performance. The schema is there; '
    'the load testing is not.'
)
para(
    'The main scalability bottleneck is external-API dependency. Rust and Axum keep our own '
    'overhead negligible, but we can\u2019t control how fast Google or Gemini responds. Redis '
    'helps only when queries repeat. For high traffic, we\u2019d need rate limiting, queue-based '
    'background enrichment, and probably a CDN or edge cache in front of the API layer.'
)

# ── Conclusion ──
heading('Conclusion and Future Work', level=1)
para(
    'PULSE demonstrates that you can build a useful local-discovery tool by combining live '
    'venue data, location awareness, LLM-generated summaries, aggressive caching, and a '
    'social-graph schema\u2014all without a massive infrastructure budget. The Rust backend '
    'stays lean, Redis makes repeated searches feel instant, Haversine provides a cheap '
    'proximity signal, and Gemini turns raw data into readable recommendations.'
)
para(
    'The parts that still need work are equally clear. External API latency dominates fresh '
    'requests. GPS accuracy drops indoors. Gemini can produce weak summaries when the input '
    'data is sparse. The social graph is designed but not yet used in ranking. The current '
    'ranking logic is mostly rule-based and doesn\u2019t learn from user behavior.'
)
para(
    'Future work should add real event feeds, stronger friend-activity signals, '
    'personalization based on past behavior, and\u2014most importantly\u2014testing with real users. '
    'We also want to collect explicit feedback (clicks, saves, shares, visit confirmations) '
    'and use it to build learned ranking models. On the infrastructure side, the deployment '
    'should move toward Kubernetes and autoscaling once traffic justifies it, and native '
    'mobile builds would improve GPS quality and enable push notifications. At a larger scale, '
    'the logged request data can support city-level trend analysis, peak-demand prediction, '
    'and smarter cache warming.'
)

# ── Contributions ──
heading('Individual Team Member Contributions', level=1)
para(
    'Mouhamed Ndiaye served as the Rust full-stack lead. He built the core backend pipeline '
    'in Rust/Axum: API route design, request validation, Google Places integration, Haversine '
    'distance computation, JSON response formatting, and the bridge between backend results '
    'and the Flutter UI. He also ran the latency experiments, compared cached vs.\u00a0uncached '
    'request performance, and cleaned venue data for consistent frontend rendering. On the '
    'analytics side, Mouhamed defined the per-request log schema (intent, result count, '
    'distance values, response time, cache status, selected card) and designed how those '
    'fields connect to product-improvement metrics.'
)
para(
    'Itayi Penda served as the DevOps and systems lead. He designed the Docker Compose setup, '
    'managed environment variables across services, configured Redis and PostgreSQL, planned '
    'the deployment pipeline, organized logging, and stress-tested service reliability. Itayi '
    'also owned the big-data analytics planning: how request logs, cache logs, and social-graph '
    'data should be stored and processed as the system scales. This included log-retention '
    'policy, batch-analysis design, dashboard wireframes, and monitoring strategy. His work '
    'bridges DevOps and analytics\u2014a big-data system needs reliable storage and clean services '
    'before it can produce useful insights.'
)
para(
    'Both team members collaborated on the project concept, system design, testing, results '
    'discussion, and final presentation. Mouhamed focused more on application code, backend '
    'logic, and analytics instrumentation. Itayi focused more on deployment, reliability, '
    'infrastructure, and big-data planning.'
)

# ── References ──
heading('References', level=1)
refs = [
    'Codebase \u2014 github.com/MouhamedN96/PULSE (branch: main)',
    'Devlin, J., Chang, M.-W., Lee, K., & Toutanova, K. (2019). BERT: Pre-training of deep bidirectional transformers for language understanding. Proc. NAACL.',
    'Google AI for Developers. (2026). Gemini API documentation.',
    'Google Maps Platform. (2026). Places API documentation.',
    'J\u00e4rvelin, K., & Kek\u00e4l\u00e4inen, J. (2002). Cumulated gain-based evaluation of IR techniques. ACM TOIS, 20(4), 422\u2013446.',
    'Kleppmann, M. (2017). Designing Data-Intensive Applications. O\u2019Reilly Media.',
    'Koren, Y., Bell, R., & Volinsky, C. (2009). Matrix factorization techniques for recommender systems. Computer, 42(8), 30\u201337.',
    'Liu, Y., Pham, T. A. N., Cong, G., & Yuan, Q. (2017). An experimental evaluation of POI recommendation in LBSNs. Proc. VLDB, 10(10), 1010\u20131021.',
    'Redis. (2026). Redis documentation.',
    'Reimers, N., & Gurevych, I. (2019). Sentence-BERT: Sentence embeddings using Siamese BERT-Networks. Proc. EMNLP.',
    'Rendle, S., Freudenthaler, C., Gantner, Z., & Schmidt-Thieme, L. (2009). BPR: Bayesian personalized ranking from implicit feedback. Proc. UAI.',
    'Resnick, P., & Varian, H. R. (1997). Recommender systems. Communications of the ACM, 40(3), 56\u201358.',
    'Ricci, F., Rokach, L., & Shapira, B. (2022). Recommender Systems Handbook. Springer.',
    'Sinnott, R. W. (1984). Virtues of the Haversine. Sky & Telescope, 68(2), 159.',
    'Ye, M., Yin, P., Lee, W.-C., & Lee, D.-L. (2011). Exploiting geographical influence for collaborative POI recommendation. Proc. SIGIR, 325\u2013334.',
    'Axum. (2026). Axum Rust web framework documentation.',
    'Docker. (2026). Docker Compose documentation.',
    'Flutter. (2026). Flutter web application documentation.',
]

for ref in refs:
    p = doc.add_paragraph(ref, style='List Bullet')
    p.paragraph_format.space_after = Pt(2)

# ── Save ──
output_path = r'c:\Users\momo-\Downloads\Kimi_Agent_STROLL Voice‑Curated Social App\stroll-rust-flutter\PULSE_FINAL_PAPER_v2.docx'
doc.save(output_path)
print(f'Saved to: {output_path}')
