import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

def add_heading(text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Arial'
        run.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
    return h

def add_para(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(11)
    return p

# Cover Page
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('PULSE: Voice-Curated Social Activity Discovery')
run.bold = True
run.font.size = Pt(16)
run.font.name = 'Arial'

doc.add_paragraph()
add_para('By: Mouhamed Ndiaye and Itayi Penda').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

# Abstract
add_heading('Abstract:', 2)
add_para('The PULSE app is designed to help bridge the gap between people looking for something to do and the overwhelming amount of disjointed local data out there. The platform will be launched as a web-based app that features a voice-curated interface to give users proactive recommendations. PULSE combines location-based Information Retrieval, real-time venue analysis, and natural language processing. The program was built using Flutter for a cross-platform front end, and Rust with Axum for a high-speed back end. For the AI integration, we initially utilized Gemini 2.0 Flash and are currently integrating a locally fine-tuned Gemma 4 E2B model via Ollama for edge-optimized summaries. Finally, the system makes use of Google Places API New, Yelp Fusion, and Eventbrite to make sure all local recommendations are based on factual, up-to-date data, backed by a Redis caching layer for speed.')

# Resources Links
add_heading('Resources links', 2)
add_para('The GGUF model and LoRA adapters will be accessible at: [HuggingFace Link TBA]')
add_para('The Github repo for the codebase along with the data used in the experiment is at: https://github.com/MouhamedN96/PULSE')
add_para('The Colab notebook for the QLoRA fine-tuning can be accessed at: [Colab Link in Repo]')

# Introduction
add_heading('Introduction:', 2)
add_para('The PULSE app aims to tackle the issue of decision fatigue and information overload when people are trying to find local activities. For your average person, actually finding a good spot—like a quiet cafe to work from or a fun rooftop bar—usually involves manually cross-referencing Google Maps, Yelp, and text threads. Additionally, most normal search engines expect exact keywords, while AI solutions tend to hallucinate places that don\'t exist or fall back on outdated hours. By creating a specialized IR system that parses natural language intent and grounds it in live API data, this app seeks to give a local assistant that delivers up-to-date, distance-ranked information with concise explanations.')

# Related Work
add_heading('Related Work', 2)
add_para('The architecture of PULSE is based on a number of different milestones in Recommender Systems, Point-of-Interest (POI) tracking, and Natural Language Processing. Retrieval-Augmented Generation is basically the framework we adapted; instead of querying a static document database, our LLM looks up facts from live external venue APIs before it generates an answer. This helps to significantly reduce hallucinations since its responses are now grounded on specific, live venue details instead of just its training data.')
add_para('We also leaned heavily into POI recommendation research, which shows that geography is a massive factor. We utilize the Haversine formula right in the Rust backend to compute distances on the fly, ensuring that even if a query is vague, the results are physically relevant. The transformer model is the engine behind the summaries we provide. We started with Gemini but are shifting towards an edge-optimized Gemma 4 E2B model. This allows the system to understand the relationships between a user\'s mood and venue tags, processing large chunks of context quickly. We also made use of Redis as our in-memory data store, which basically acts as the short-term memory of the system, caching repeated nearby searches to drastically reduce latency and API costs. Finally, we set up PostgreSQL to act as the long-term memory for a social graph, allowing future recalls of what places friends have saved or visited.')

# Datasets
add_heading('Base Datasets', 2)
add_para('Datasets:')
add_para('Google Places API New: ', bold_prefix='').runs[0].text = '- Google Places API New: This is the main source of live venue data, providing ratings, hours, photos, and coordinates.'
add_para('Yelp Fusion API: ', bold_prefix='').runs[0].text = '- Yelp Fusion API: Chosen for detailed business categories and price range data to handle affordability queries.'
add_para('Eventbrite API: ', bold_prefix='').runs[0].text = '- Eventbrite API: Adds a time-based dimension for live events happening right now or this weekend.'
add_para('nyc_seed.json (Custom Training Data): ', bold_prefix='').runs[0].text = '- nyc_seed.json (Custom Training Data): A hand-curated dataset of NYC venues we used to generate Q&A pairs for our custom Gemma fine-tuning.'

add_para('Collections Method: ', bold_prefix='').runs[0].text = 'Collections Method: We collect data at runtime when a user makes a request. The backend builds a cache key; if there\'s no hit, it reaches out to the APIs. For our local model fine-tuning, we scraped a subset of venues into the `nyc_seed.json` file and built a script to generate golden queries (e.g., "best tapas near me") and ideal summary targets.'

add_para('Issues Encountered: ', bold_prefix='').runs[0].text = 'Issues Encountered: Venue records often had missing fields like hours or photos, which we mitigated by enforcing safe defaults in Rust. For the AI, we noticed cold-start delays with cloud models, which pushed us towards building a local, edge-optimized GGUF model via Ollama to keep inference fast and private.'

# Detailed Approach
add_heading('Detailed Approach:', 2)
add_para('The method we used was a multi-layered architecture meant for speed, accuracy, and eventual privacy when it came to local data. The "Data Layer" relies on Redis to act as a high-speed cache for venue records, while PostgreSQL handles user accounts and a planned social graph. Our backend pipeline is written in Rust using the Axum framework, which basically takes the user\'s coordinates and query, hits the cache or external APIs, calculates the Haversine distance for sorting, and then prepares the data for the AI.')
add_para('The Intelligence Layer of the app works as its brain. Initially, we used Gemini to summarize the venue data. However, to enhance privacy and reduce latency, we are migrating to an edge-optimized setup using Ollama running a fine-tuned Gemma 4 E2B model. When a question is received, the system doesn\'t just dump raw data; it pulls the top venues, injects them into the LLM\'s prompt, and generates a two-sentence, practical summary explaining *why* the place fits the user\'s mood. After this, the Delivery layer uses a Flutter Progressive Web App (PWA) to give a fluid, app-like interface with swipeable cards and voice input.')

add_para('To finetune our Gemma model, we set up a QLoRA pipeline in a Colab notebook. The choices behind this setup:')
add_para('- We utilized `pulse-gemma-e2b-q4.gguf` as our target format to ensure it runs efficiently on consumer hardware via Ollama.')
add_para('- The training set was generated from our `nyc_seed.json`, creating triplets of (User Query, Top Venues, Ideal Summary).')
add_para('- We formatted the data into JSONL and used the HuggingFace `trl` library with 4-bit quantization to keep memory usage low on a T4 GPU.')

# Experiments/Results
add_heading('Experimentation Results & Discussions:', 2)
add_para('- Experimental Setup:')
add_para('To mirror real-world usage, we tested the system on common queries like "quiet cafe for work", "fun rooftop bar", and "cheap eats". We compared fresh API calls against Redis cache hits, and raw venue lists against AI-summarized cards. For the AI transition, we set up a Weights & Biases (W&B) integration in our Colab notebook to monitor the QLoRA fine-tuning of the Gemma 4 E2B model.')

add_para('- System Architecture and Toolkit:')
add_para('We utilized the Flutter framework for Cross-platform UI (web/mobile). On the server side, we used Rust Axum for high concurrency and low memory overhead. Persistence and caching are handled by PostgreSQL and Redis. For inference, the `recommendation_api` in Rust handles provider dispatch, switching between standard cloud APIs and our local Ollama instance running the GGUF model, avoiding cross-backend collisions.')

add_para('- Evaluation methodology')
add_para('We went for a Quantitative evaluation of latency (comparing cached vs uncached requests) and training metrics for our model. For the model, train and eval loss were logged to Weights & Biases during the SFTTrainer run.')

add_para('For the backend performance:')
add_para('- Cache hits consistently returned in under 80ms, proving Redis is essential for the mobile experience.')
add_para('- Fresh calls took noticeably longer due to external API latency, highlighting the need for local AI inference to cut down on at least one external hop.')

add_para('- Feed Recommendation Evaluation & Backend Experiments:')
add_para('We evaluated our retrieval and ranking system using offline and online metrics. Our primary offline metric is NDCG@5, which measures the position-weighted quality of the top 5 results returned to the user, penalizing highly relevant venues that are ranked too low. The client-side flutter app performs a re-ranking based on a composite score considering rating (24%), query affinity (14%), open status (14%), user save history (14%), review count (12%), and inferred category/tag preferences.')
add_para('In our benchmark experiments against Google Places (Run 001) and Yelp Fusion APIs (Run 005), we found that our client-side re-ranking consistently improved overall NDCG@5 over the default API orders (e.g., +0.027 improvement for Places and +0.0305 for Yelp). For instance, in our Yelp run, client NDCG@5 reached 0.9322. We also implemented a local SQLite mock fallback using curated seed data (nyc_seed.json) achieving high retrieval coverage (15/15) without relying on live API calls.')

add_para('For the Gemma 4 E2B Fine-Tuning (Ongoing):')
add_para('- We successfully configured the Colab pipeline to handle the JSONL venue data.')
add_para('- The QLoRA setup targeted the q_proj, k_proj, v_proj, and o_proj modules with an r-value of 16.')
add_para('- We resolved data-loading issues in the notebook by moving to a more stable Google Drive/Sidebar upload approach, ensuring the training run doesn\'t crash due to missing seed files.')
add_para('- Training metrics are currently streaming to the W&B dashboard under the `pulse-gemma-finetune` project.')

# Conclusion
add_heading('Conclusion:', 2)
add_para('PULSE shows that local social activity discovery can be vastly improved by combining live place data, location awareness, aggressive caching, and generative AI. We successfully built a pipeline that captures user intent and returns useful, summarized cards. The Rust backend and Redis caching proved to be a highly effective combination for speed.')
add_para('The transition towards an edge-optimized Gemma 4 E2B model via Ollama represents a significant step in reducing external API dependency and improving user privacy. While the fine-tuning process requires careful data formatting and environment management (as seen in our Colab debugging), the resulting GGUF model will allow the application to generate context-aware summaries entirely locally. Future work will focus on integrating the social graph for friend-based recommendations and expanding the analytics layer to monitor cache hit rates and popular categories at scale.')

# Sources
add_heading('Sources:', 2)
add_para('- Codebase: https://github.com/MouhamedN96/PULSE')
add_para('- Recommender Systems Handbook (Ricci et al., 2022)')
add_para('- Exploiting geographical influence for collaborative POI recommendation (Ye et al., 2011)')
add_para('- Google Maps Platform. (2026). Places API documentation.')
add_para('- Redis. (2026). Redis documentation.')
add_para('- Axum. (2026). Axum Rust web framework documentation.')
add_para('- Flutter. (2026). Flutter web application documentation.')
add_para('- HuggingFace PEFT & TRL documentation for QLoRA fine-tuning.')

output_path = r'c:\Users\momo-\Downloads\Kimi_Agent_STROLL Voice‑Curated Social App\stroll-rust-flutter\PULSE_FINAL_PAPER_v4_IR_Style.docx'
doc.save(output_path)
print('Saved successfully.')
